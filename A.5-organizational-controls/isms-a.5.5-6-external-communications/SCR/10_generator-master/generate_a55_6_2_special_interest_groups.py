#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# =============================================================================
# SPDX-License-Identifier: AGPL-3.0-or-later OR LicenseRef-ISMS-Commercial
# Copyright (c) 2025-2026 ISMS Core Contributors
# =============================================================================
"""
================================================================================
ISMS-IMP-A.5.5-6.S2 - Special Interest Groups Register
================================================================================

ISO/IEC 27001:2022 Controls A.5.5 & A.5.6: External Communications
Assessment Domain 2 of 4: Special Interest Groups Register

This script generates a comprehensive Excel workbook for managing contacts with
special interest groups including ISACs, security forums, professional associations,
and vendor communities as required by ISO 27001:2022 Control A.5.6.
================================================================================
"""

import logging
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# =============================================================================
# LOGGING CONFIGURATION
# =============================================================================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

# =============================================================================
# DOCUMENT METADATA
# =============================================================================
DOCUMENT_ID = "ISMS-IMP-A.5.5-6.S2"
WORKBOOK_NAME = "Special Interest Groups Register"
CONTROL_ID = "A.5.5-6"
CONTROL_NAME = "Contact with Authorities & Special Interest Groups"
CONTROL_REF = f"ISO/IEC 27001:2022 - Controls A.5.5 & A.5.6: {CONTROL_NAME}"

# Timestamps
GENERATED_DATE = datetime.now().strftime("%d.%m.%Y")
GENERATED_TIMESTAMP = datetime.now().strftime("%Y%m%d")

# Output filename
OUTPUT_FILENAME = f"{DOCUMENT_ID}_{WORKBOOK_NAME.replace(' ', '_')}_{GENERATED_TIMESTAMP}.xlsx"

# =============================================================================
# STYLING CONSTANTS
# =============================================================================
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

INPUT_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


# =============================================================================
# WORKBOOK GENERATION FUNCTIONS
# =============================================================================

def create_instructions_sheet(ws):
    """Create the Instructions sheet."""
    ws.title = "Instructions"

    instructions = [
        ["ISMS-IMP-A.5.5-6.S2 - Special Interest Groups Register"],
        [""],
        ["PURPOSE"],
        ["This workbook maintains a register of special interest groups, professional forums,"],
        ["and information sharing communities as required by ISO 27001:2022 Control A.5.6."],
        [""],
        ["ISO 27001:2022 CONTROL A.5.6 - CONTACT WITH SPECIAL INTEREST GROUPS"],
        ["Appropriate contacts with special interest groups or other specialist security forums"],
        ["and professional associations should be maintained."],
        [""],
        ["SHEETS"],
        ["1. Instructions - This guidance sheet"],
        ["2. Groups_Registry - Master list of special interest groups"],
        ["3. Membership_Details - Membership status and benefits tracking"],
        ["4. Engagement_Log - Record of engagement activities"],
        ["5. Intelligence_Received - Threat intelligence and advisories received"],
        ["6. Contribution_Log - Our contributions to groups"],
        ["7. Evidence_Register - Evidence documentation for audits"],
        ["8. Approval_SignOff - Management approval workflow"],
        [""],
        ["GROUP CATEGORIES"],
        ["- ISACs (Information Sharing and Analysis Centers)"],
        ["- Security Forums (FIRST, (ISC)2, ISACA chapters)"],
        ["- Professional Associations (ACM, IEEE, local IT associations)"],
        ["- Vendor Security Communities (Microsoft, AWS, Google security groups)"],
        ["- Standards Bodies Participation (ISO working groups)"],
        ["- Threat Intelligence Sharing (MISP communities, sector groups)"],
        ["- Open Source Security Communities (OWASP, Linux Foundation)"],
        ["- Academic/Research Partnerships"],
        [""],
        ["BENEFITS OF PARTICIPATION"],
        ["- Early warning of emerging threats"],
        ["- Access to threat intelligence and IOCs"],
        ["- Best practice sharing and peer learning"],
        ["- Professional development opportunities"],
        ["- Networking and relationship building"],
        ["- Standards development influence"],
        [""],
        ["DATA VALIDATION"],
        ["- Yellow cells are input fields"],
        ["- Dropdown lists enforce valid values"],
        ["- Group_ID must be unique (SIG-XXXX format)"],
        [""],
        [f"Generated: {GENERATED_DATE}"],
        [f"Control Reference: {CONTROL_REF}"],
    ]

    for row_num, row_data in enumerate(instructions, 1):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            if row_num == 1:
                cell.font = Font(bold=True, size=14)
            elif value in ["PURPOSE", "SHEETS", "GROUP CATEGORIES", "BENEFITS OF PARTICIPATION", "DATA VALIDATION",
                          "ISO 27001:2022 CONTROL A.5.6 - CONTACT WITH SPECIAL INTEREST GROUPS"]:
                cell.font = Font(bold=True, size=11)

    ws.column_dimensions['A'].width = 90


def create_groups_registry_sheet(ws):
    """Create the Groups Registry sheet."""
    ws.title = "Groups_Registry"

    headers = [
        "Group_ID", "Group_Name", "Group_Type", "Focus_Area",
        "Geographic_Scope", "Website", "Primary_Contact",
        "Contact_Email", "Internal_Owner", "Owner_Department",
        "Membership_Status", "Member_Since", "Membership_Level",
        "Annual_Cost", "Value_Rating", "Last_Engagement", "Next_Review",
        "Strategic_Priority", "Notes"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    type_dv = DataValidation(
        type="list",
        formula1='"ISAC,Security Forum,Professional Association,Vendor Community,Standards Body,Threat Intel Sharing,Open Source Community,Academic Partnership,Other"'
    )
    ws.add_data_validation(type_dv)
    type_dv.add('C2:C100')

    focus_dv = DataValidation(
        type="list",
        formula1='"Cybersecurity General,Financial Services,Healthcare,Critical Infrastructure,Cloud Security,Application Security,Threat Intelligence,Governance/Compliance,Technical Standards,Other"'
    )
    ws.add_data_validation(focus_dv)
    focus_dv.add('D2:D100')

    scope_dv = DataValidation(
        type="list",
        formula1='"Global,Europe,DACH Region,Switzerland,United States,Asia-Pacific,Other"'
    )
    ws.add_data_validation(scope_dv)
    scope_dv.add('E2:E100')

    status_dv = DataValidation(
        type="list",
        formula1='"Active Member,Pending,Suspended,Former Member,Under Evaluation"'
    )
    ws.add_data_validation(status_dv)
    status_dv.add('K2:K100')

    level_dv = DataValidation(
        type="list",
        formula1='"Full,Associate,Observer,Sponsor,Free Tier"'
    )
    ws.add_data_validation(level_dv)
    level_dv.add('M2:M100')

    value_dv = DataValidation(
        type="list",
        formula1='"High,Medium,Low,Under Review"'
    )
    ws.add_data_validation(value_dv)
    value_dv.add('O2:O100')

    priority_dv = DataValidation(
        type="list",
        formula1='"Critical,High,Medium,Low"'
    )
    ws.add_data_validation(priority_dv)
    priority_dv.add('R2:R100')

    # Format input rows
    for row in range(2, 31):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [12, 35, 22, 25, 15, 35, 25, 30, 20, 18, 15, 15, 15, 12, 12, 15, 15, 15, 35]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def create_membership_details_sheet(ws):
    """Create the Membership Details sheet."""
    ws.title = "Membership_Details"

    headers = [
        "Group_ID", "Group_Name", "Membership_Start", "Membership_End",
        "Membership_Type", "Access_Level", "Portal_Credentials",
        "Mailing_Lists_Subscribed", "Events_Access", "Intel_Feed_Access",
        "Publication_Access", "Voting_Rights", "Annual_Fee",
        "Billing_Contact", "Contract_Reference", "Benefits_Summary"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    type_dv = DataValidation(
        type="list",
        formula1='"Corporate,Individual,Government,Academic,Non-Profit"'
    )
    ws.add_data_validation(type_dv)
    type_dv.add('E2:E100')

    access_dv = DataValidation(
        type="list",
        formula1='"Full,Restricted,Read-Only,None"'
    )
    ws.add_data_validation(access_dv)
    access_dv.add('F2:F100')

    yn_dv = DataValidation(type="list", formula1='"Yes,No,Limited"')
    ws.add_data_validation(yn_dv)
    yn_dv.add('I2:I100')
    yn_dv.add('J2:J100')
    yn_dv.add('K2:K100')
    yn_dv.add('L2:L100')

    # Format input rows
    for row in range(2, 31):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [12, 30, 15, 15, 15, 15, 18, 30, 15, 15, 15, 15, 12, 25, 25, 40]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def create_engagement_log_sheet(ws):
    """Create the Engagement Log sheet."""
    ws.title = "Engagement_Log"

    headers = [
        "Engagement_ID", "Group_ID", "Group_Name", "Date",
        "Engagement_Type", "Description", "Our_Representative",
        "Key_Topics", "Outcomes", "Action_Items",
        "Follow_Up_Date", "Evidence_Reference", "Notes"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    type_dv = DataValidation(
        type="list",
        formula1='"Meeting,Webinar,Conference,Working Group,Training,Intel Briefing,Publication Review,Networking Event,Other"'
    )
    ws.add_data_validation(type_dv)
    type_dv.add('E2:E200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [15, 12, 30, 15, 20, 45, 25, 40, 40, 35, 15, 25, 35]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def create_intelligence_received_sheet(ws):
    """Create the Intelligence Received sheet."""
    ws.title = "Intelligence_Received"

    headers = [
        "Intel_ID", "Group_ID", "Group_Name", "Date_Received",
        "Intel_Type", "Classification", "Title", "Summary",
        "Relevance_To_Org", "Action_Taken", "Distributed_To",
        "Distribution_Date", "Evidence_Location", "Notes"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    type_dv = DataValidation(
        type="list",
        formula1='"Threat Advisory,IOC Feed,Vulnerability Alert,Best Practice,Research Report,Incident Report,Policy Update,Standards Update,Other"'
    )
    ws.add_data_validation(type_dv)
    type_dv.add('E2:E200')

    class_dv = DataValidation(
        type="list",
        formula1='"TLP:RED,TLP:AMBER,TLP:GREEN,TLP:CLEAR,Internal Only,Public"'
    )
    ws.add_data_validation(class_dv)
    class_dv.add('F2:F200')

    relevance_dv = DataValidation(
        type="list",
        formula1='"Critical,High,Medium,Low,Informational"'
    )
    ws.add_data_validation(relevance_dv)
    relevance_dv.add('I2:I200')

    # Format input rows
    for row in range(2, 101):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [12, 12, 25, 15, 20, 15, 40, 50, 15, 35, 25, 15, 30, 35]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def create_contribution_log_sheet(ws):
    """Create the Contribution Log sheet."""
    ws.title = "Contribution_Log"

    headers = [
        "Contribution_ID", "Group_ID", "Group_Name", "Date",
        "Contribution_Type", "Title", "Description",
        "Contributor", "Approval_Status", "Approved_By",
        "Classification_Check", "Publication_Status", "Notes"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    type_dv = DataValidation(
        type="list",
        formula1='"IOC Sharing,Incident Report,Best Practice,Research Paper,Presentation,Working Group Contribution,Standards Input,Tool/Script,Other"'
    )
    ws.add_data_validation(type_dv)
    type_dv.add('E2:E100')

    approval_dv = DataValidation(
        type="list",
        formula1='"Pending,Approved,Rejected,N/A"'
    )
    ws.add_data_validation(approval_dv)
    approval_dv.add('I2:I100')

    class_dv = DataValidation(
        type="list",
        formula1='"Passed,Failed,Pending,N/A"'
    )
    ws.add_data_validation(class_dv)
    class_dv.add('K2:K100')

    pub_dv = DataValidation(
        type="list",
        formula1='"Published,Pending,Withdrawn"'
    )
    ws.add_data_validation(pub_dv)
    pub_dv.add('L2:L100')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [15, 12, 25, 15, 25, 35, 45, 25, 15, 20, 18, 18, 35]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def create_evidence_register_sheet(ws):
    """Create the Evidence Register sheet."""
    ws.title = "Evidence_Register"

    headers = [
        "Evidence_ID", "Evidence_Type", "Description", "Related_Group_ID",
        "Date_Created", "Created_By", "Storage_Location",
        "Retention_Period", "Review_Date", "Status", "Notes"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    type_dv = DataValidation(
        type="list",
        formula1='"Membership Certificate,Meeting Minutes,Intel Report,Contribution Record,Training Certificate,Event Attendance,Other"'
    )
    ws.add_data_validation(type_dv)
    type_dv.add('B2:B100')

    status_dv = DataValidation(type="list", formula1='"Current,Archived,Pending Review"')
    ws.add_data_validation(status_dv)
    status_dv.add('J2:J100')

    # Format input rows
    for row in range(2, 31):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [12, 22, 40, 18, 15, 20, 40, 15, 15, 15, 35]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def create_approval_signoff_sheet(ws):
    """Create the Approval Sign-Off sheet."""
    ws.title = "Approval_SignOff"

    headers = [
        "Approval_ID", "Review_Period", "Review_Date", "Reviewer_Name",
        "Reviewer_Role", "Registry_Complete", "Memberships_Active",
        "Value_Assessed", "Approval_Status", "Signature_Date",
        "Next_Review_Date", "Comments"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data validations
    role_dv = DataValidation(
        type="list",
        formula1='"CISO,Security Manager,IT Director,Risk Manager,Compliance Officer"'
    )
    ws.add_data_validation(role_dv)
    role_dv.add('E2:E20')

    yn_dv = DataValidation(type="list", formula1='"Yes,No,Partial"')
    ws.add_data_validation(yn_dv)
    yn_dv.add('F2:F20')
    yn_dv.add('G2:G20')
    yn_dv.add('H2:H20')

    status_dv = DataValidation(type="list", formula1='"Approved,Rejected,Pending,Conditional"')
    ws.add_data_validation(status_dv)
    status_dv.add('I2:I20')

    # Format input rows
    for row in range(2, 11):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

    # Column widths
    widths = [12, 15, 15, 25, 20, 18, 18, 15, 15, 15, 18, 40]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def generate_workbook():
    """Generate the complete workbook."""
    logger.info("=" * 70)
    logger.info(f"Generating {DOCUMENT_ID} - {WORKBOOK_NAME}")
    logger.info("=" * 70)

    wb = Workbook()

    # Create all sheets
    create_instructions_sheet(wb.active)
    create_groups_registry_sheet(wb.create_sheet())
    create_membership_details_sheet(wb.create_sheet())
    create_engagement_log_sheet(wb.create_sheet())
    create_intelligence_received_sheet(wb.create_sheet())
    create_contribution_log_sheet(wb.create_sheet())
    create_evidence_register_sheet(wb.create_sheet())
    create_approval_signoff_sheet(wb.create_sheet())

    # Save workbook
    wb.save(OUTPUT_FILENAME)
    logger.info(f"Workbook saved: {OUTPUT_FILENAME}")
    logger.info("=" * 70)

    return OUTPUT_FILENAME


if __name__ == "__main__":
    generate_workbook()

# =============================================================================
# QA_VERIFIED: 2026-02-01
# QA_STATUS: PASSED
# QA_TOOL: Claude Code
# CHANGES: Initial creation for A.5.5-6 External Communications control
# =============================================================================

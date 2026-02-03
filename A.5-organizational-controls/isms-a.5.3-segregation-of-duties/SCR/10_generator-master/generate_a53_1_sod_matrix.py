#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# =============================================================================
# SPDX-License-Identifier: AGPL-3.0-or-later OR LicenseRef-ISMS-Commercial
# Copyright (c) 2025-2026 ISMS Core Contributors
# =============================================================================
"""
================================================================================
ISMS-IMP-A.5.3.1 - SoD Matrix Assessment
================================================================================

ISO/IEC 27001:2022 Control A.5.3: Segregation of Duties
Assessment Workbook 1 of 4: SoD Matrix Assessment

This script generates an Excel workbook for documenting and tracking
segregation of duties conflicts, role inventories, and remediation efforts.
================================================================================
"""

import logging
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# =============================================================================
# LOGGING CONFIGURATION
# =============================================================================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

# =============================================================================
# DOCUMENT METADATA
# =============================================================================
DOCUMENT_ID = "ISMS-IMP-A.5.3.1"
WORKBOOK_NAME = "SoD Matrix Assessment"
CONTROL_ID = "A.5.3"
CONTROL_NAME = "Segregation of Duties"
CONTROL_REF = f"ISO/IEC 27001:2022 - Control {CONTROL_ID}: {CONTROL_NAME}"

# Timestamps
GENERATED_DATE = datetime.now().strftime("%d.%m.%Y")
GENERATED_TIMESTAMP = datetime.now().strftime("%Y%m%d")

# Output filename
OUTPUT_FILENAME = f"{DOCUMENT_ID}_{WORKBOOK_NAME.replace(' ', '_')}_{GENERATED_TIMESTAMP}.xlsx"

# =============================================================================
# STYLING CONSTANTS
# =============================================================================
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
TITLE_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

SUBHEADER_FILL = PatternFill(start_color="D6DCE4", end_color="D6DCE4", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, size=10)

INPUT_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
LOCKED_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

# Conflict type colors
CONFLICT_X_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
CONFLICT_C_FILL = PatternFill(start_color="FABF8F", end_color="FABF8F", fill_type="solid")
CONFLICT_M_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# =============================================================================
# DATA CONSTANTS
# =============================================================================
PROCESS_DOMAINS = [
    "Financial",
    "IT Operations",
    "HR",
    "Procurement",
    "Security",
    "Change Management",
    "Other"
]

RISK_LEVELS = ["Critical", "High", "Medium", "Low"]

CONFLICT_TYPES = ["X", "C", "M", "-"]

GAP_STATUSES = ["Open", "Mitigated", "Resolved", "Accepted"]

REMEDIATION_TYPES = [
    "Role Removal",
    "Role Reassignment",
    "Process Redesign",
    "Compensating Control"
]

REMEDIATION_STATUSES = ["Not Started", "In Progress", "Completed", "Cancelled"]

EXCEPTION_STATUSES = ["Active", "Expired", "Revoked"]

REVIEW_FREQUENCIES = ["Monthly", "Quarterly", "Semi-Annual", "Annual"]

DEPARTMENTS = [
    "Executive",
    "Finance",
    "IT",
    "Operations",
    "HR",
    "Legal",
    "Sales",
    "Marketing",
    "Engineering",
    "Support",
    "Procurement",
    "Security"
]


# =============================================================================
# WORKBOOK GENERATION FUNCTIONS
# =============================================================================

def apply_header_style(cell):
    """Apply standard header styling to a cell."""
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGNMENT
    cell.border = THIN_BORDER


def apply_input_style(cell):
    """Apply input cell styling."""
    cell.fill = INPUT_FILL
    cell.border = THIN_BORDER


def create_instructions_sheet(ws):
    """Create the Instructions sheet."""
    ws.title = "Instructions"

    instructions = [
        [f"{DOCUMENT_ID} - {WORKBOOK_NAME}"],
        [""],
        ["PURPOSE"],
        ["This workbook documents segregation of duties (SoD) assessment per ISMS-POL-A.5.3."],
        ["It identifies roles, defines conflicts, tracks current assignments, and manages remediation."],
        [""],
        ["SHEETS"],
        ["1. Instructions - This guidance sheet"],
        ["2. Role_Inventory - Catalogue of all roles in scope"],
        ["3. Conflict_Matrix - Role-to-role conflict definitions"],
        ["4. Current_Assignments - Actual role assignments by person"],
        ["5. Gap_Analysis - Identified conflicts requiring action"],
        ["6. Remediation_Tracker - Actions to resolve gaps"],
        ["7. Exception_Register - Approved exceptions with compensating controls"],
        ["8. Approval_SignOff - Review and approval signatures"],
        [""],
        ["WORKFLOW"],
        ["1. Populate Role_Inventory with all organisational roles"],
        ["2. Define conflicts in Conflict_Matrix (X=Hard, C=Conditional, M=Monitor)"],
        ["3. Document current assignments in Current_Assignments"],
        ["4. Analyse gaps by comparing assignments against conflict matrix"],
        ["5. Create remediation plans or document exceptions"],
        ["6. Obtain approval signatures"],
        [""],
        ["CONFLICT CLASSIFICATIONS"],
        ["X - Hard Conflict: Must never be combined; no exceptions permitted"],
        ["C - Conditional Conflict: Requires compensating controls if combined"],
        ["M - Monitoring Required: May be combined with enhanced monitoring"],
        ["- - No Conflict: Roles may be combined without restriction"],
        [""],
        [f"Generated: {GENERATED_DATE}"],
        [f"Control Reference: {CONTROL_REF}"],
    ]

    for row_num, row_data in enumerate(instructions, 1):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            if row_num == 1:
                cell.font = Font(bold=True, size=14)
                cell.fill = TITLE_FILL
                cell.font = TITLE_FONT
            elif value in ["PURPOSE", "SHEETS", "WORKFLOW", "CONFLICT CLASSIFICATIONS"]:
                cell.font = Font(bold=True, size=11)

    ws.column_dimensions['A'].width = 80


def create_role_inventory_sheet(ws):
    """Create the Role Inventory sheet."""
    ws.title = "Role_Inventory"

    headers = [
        "Role_ID", "Role_Name", "Department", "Process_Domain", "Risk_Level",
        "Description", "Key_Duties", "System_Access", "Active"
    ]

    widths = [15, 30, 20, 20, 12, 40, 40, 30, 10]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    dept_dv = DataValidation(type="list", formula1=f'"{",".join(DEPARTMENTS)}"')
    ws.add_data_validation(dept_dv)
    dept_dv.add('C2:C200')

    domain_dv = DataValidation(type="list", formula1=f'"{",".join(PROCESS_DOMAINS)}"')
    ws.add_data_validation(domain_dv)
    domain_dv.add('D2:D200')

    risk_dv = DataValidation(type="list", formula1=f'"{",".join(RISK_LEVELS)}"')
    ws.add_data_validation(risk_dv)
    risk_dv.add('E2:E200')

    active_dv = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(active_dv)
    active_dv.add('I2:I200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Add sample roles
    sample_roles = [
        ("ROLE-FIN-001", "AP Clerk", "Finance", "Financial", "Medium",
         "Processes vendor invoices", "Invoice entry, payment prep", "SAP, Treasury", "Yes"),
        ("ROLE-FIN-002", "AP Manager", "Finance", "Financial", "High",
         "Manages AP team and approvals", "Approve payments, supervise", "SAP, Treasury", "Yes"),
        ("ROLE-FIN-003", "Treasurer", "Finance", "Financial", "Critical",
         "Manages treasury operations", "Execute payments, banking", "Treasury, Banking", "Yes"),
        ("ROLE-IT-001", "Developer", "IT", "IT Operations", "Medium",
         "Develops application code", "Code, unit test", "Git, IDE, Dev DB", "Yes"),
        ("ROLE-IT-002", "Release Manager", "IT", "Change Management", "High",
         "Manages production deployments", "Deploy, release approval", "CI/CD, Prod Access", "Yes"),
    ]

    for row_idx, role_data in enumerate(sample_roles, 2):
        for col_idx, value in enumerate(role_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_conflict_matrix_sheet(ws):
    """Create the Conflict Matrix sheet."""
    ws.title = "Conflict_Matrix"

    # Headers - sample roles for demonstration
    sample_roles = ["ROLE-FIN-001", "ROLE-FIN-002", "ROLE-FIN-003", "ROLE-IT-001", "ROLE-IT-002"]

    # Title row
    cell = ws.cell(row=1, column=1, value="Role \\ Role")
    apply_header_style(cell)

    # Column headers (roles)
    for col, role in enumerate(sample_roles, 2):
        cell = ws.cell(row=1, column=col, value=role)
        apply_header_style(cell)

    # Row headers and conflict matrix
    conflict_data = [
        # AP Clerk conflicts
        ["ROLE-FIN-001", "-", "-", "X", "-", "-"],
        # AP Manager conflicts
        ["ROLE-FIN-002", "-", "-", "X", "-", "-"],
        # Treasurer conflicts
        ["ROLE-FIN-003", "X", "X", "-", "-", "-"],
        # Developer conflicts
        ["ROLE-IT-001", "-", "-", "-", "-", "X"],
        # Release Manager conflicts
        ["ROLE-IT-002", "-", "-", "-", "X", "-"],
    ]

    for row_idx, row_data in enumerate(conflict_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == 1:
                apply_header_style(cell)
            else:
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center")
                if value == "X":
                    cell.fill = CONFLICT_X_FILL
                    cell.font = Font(bold=True)
                elif value == "C":
                    cell.fill = CONFLICT_C_FILL
                elif value == "M":
                    cell.fill = CONFLICT_M_FILL

    # Data validation for conflict types
    conflict_dv = DataValidation(type="list", formula1=f'"{",".join(CONFLICT_TYPES)}"')
    ws.add_data_validation(conflict_dv)
    conflict_dv.add('B2:Z200')

    # Column widths
    ws.column_dimensions['A'].width = 15
    for col in range(2, 20):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # Add legend
    ws.cell(row=10, column=1, value="LEGEND:").font = Font(bold=True)
    ws.cell(row=11, column=1, value="X = Hard Conflict (never combine)")
    ws.cell(row=11, column=1).fill = CONFLICT_X_FILL
    ws.cell(row=12, column=1, value="C = Conditional (compensating controls)")
    ws.cell(row=12, column=1).fill = CONFLICT_C_FILL
    ws.cell(row=13, column=1, value="M = Monitoring Required")
    ws.cell(row=13, column=1).fill = CONFLICT_M_FILL
    ws.cell(row=14, column=1, value="- = No Conflict")


def create_current_assignments_sheet(ws):
    """Create the Current Assignments sheet."""
    ws.title = "Current_Assignments"

    headers = [
        "Person_ID", "Name", "Department", "Primary_Role", "Additional_Roles",
        "Assignment_Date", "Last_Review", "Manager", "Notes"
    ]

    widths = [12, 25, 20, 30, 50, 15, 15, 25, 30]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    dept_dv = DataValidation(type="list", formula1=f'"{",".join(DEPARTMENTS)}"')
    ws.add_data_validation(dept_dv)
    dept_dv.add('C2:C200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))


def create_gap_analysis_sheet(ws):
    """Create the Gap Analysis sheet."""
    ws.title = "Gap_Analysis"

    headers = [
        "Gap_ID", "Person_ID", "Name", "Conflicting_Roles", "Conflict_Type",
        "Risk_Level", "Identified_Date", "Status", "Notes"
    ]

    widths = [15, 12, 25, 50, 12, 12, 15, 15, 40]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    conflict_type_dv = DataValidation(type="list", formula1='"X,C,M"')
    ws.add_data_validation(conflict_type_dv)
    conflict_type_dv.add('E2:E200')

    risk_dv = DataValidation(type="list", formula1=f'"{",".join(RISK_LEVELS)}"')
    ws.add_data_validation(risk_dv)
    risk_dv.add('F2:F200')

    status_dv = DataValidation(type="list", formula1=f'"{",".join(GAP_STATUSES)}"')
    ws.add_data_validation(status_dv)
    status_dv.add('H2:H200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))


def create_remediation_tracker_sheet(ws):
    """Create the Remediation Tracker sheet."""
    ws.title = "Remediation_Tracker"

    headers = [
        "Remediation_ID", "Gap_ID", "Action_Type", "Description", "Owner",
        "Target_Date", "Status", "Completion_Date", "Evidence_Ref"
    ]

    widths = [15, 15, 18, 50, 25, 15, 15, 15, 30]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    type_dv = DataValidation(type="list", formula1=f'"{",".join(REMEDIATION_TYPES)}"')
    ws.add_data_validation(type_dv)
    type_dv.add('C2:C200')

    status_dv = DataValidation(type="list", formula1=f'"{",".join(REMEDIATION_STATUSES)}"')
    ws.add_data_validation(status_dv)
    status_dv.add('G2:G200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))


def create_exception_register_sheet(ws):
    """Create the Exception Register sheet."""
    ws.title = "Exception_Register"

    headers = [
        "Exception_ID", "Gap_ID", "Justification", "Compensating_Controls",
        "Risk_Acceptance", "Approval_Date", "Expiry_Date", "Review_Frequency",
        "Last_Review", "Status"
    ]

    widths = [15, 15, 50, 60, 25, 15, 15, 15, 15, 12]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    freq_dv = DataValidation(type="list", formula1=f'"{",".join(REVIEW_FREQUENCIES)}"')
    ws.add_data_validation(freq_dv)
    freq_dv.add('H2:H200')

    status_dv = DataValidation(type="list", formula1=f'"{",".join(EXCEPTION_STATUSES)}"')
    ws.add_data_validation(status_dv)
    status_dv.add('J2:J200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))


def create_approval_signoff_sheet(ws):
    """Create the Approval Sign-Off sheet."""
    ws.title = "Approval_SignOff"

    # Title
    cell = ws.cell(row=1, column=1, value="SoD Matrix Assessment - Approval Sign-Off")
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    ws.merge_cells('A1:E1')

    headers = ["Role", "Name", "Date", "Signature", "Comments"]
    widths = [30, 25, 15, 20, 50]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Pre-populated roles
    roles = [
        "Assessment Lead",
        "Process Owner - Financial",
        "Process Owner - IT",
        "IT Security Manager",
        "Internal Audit",
        "CISO",
        "Executive Management (Risk Acceptance)"
    ]

    for row_idx, role in enumerate(roles, 4):
        ws.cell(row=row_idx, column=1, value=role)
        for col in range(1, 6):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = THIN_BORDER
            if col > 1:
                cell.fill = INPUT_FILL


def generate_workbook():
    """Generate the complete assessment workbook."""
    logger.info(f"Generating {DOCUMENT_ID} - {WORKBOOK_NAME}")

    wb = Workbook()

    # Remove default sheet
    default_sheet = wb.active

    # Create all sheets
    create_instructions_sheet(wb.create_sheet())
    create_role_inventory_sheet(wb.create_sheet())
    create_conflict_matrix_sheet(wb.create_sheet())
    create_current_assignments_sheet(wb.create_sheet())
    create_gap_analysis_sheet(wb.create_sheet())
    create_remediation_tracker_sheet(wb.create_sheet())
    create_exception_register_sheet(wb.create_sheet())
    create_approval_signoff_sheet(wb.create_sheet())

    # Remove default sheet
    wb.remove(default_sheet)

    # Save workbook
    wb.save(OUTPUT_FILENAME)
    logger.info(f"Workbook saved: {OUTPUT_FILENAME}")

    return OUTPUT_FILENAME


# =============================================================================
# MAIN EXECUTION
# =============================================================================
if __name__ == "__main__":
    generate_workbook()

# =============================================================================
# QA_VERIFIED: 2026-02-03
# QA_STATUS: PASSED - INITIAL CREATION
# QA_TOOL: Claude Code
# CHANGES: Initial creation for A.5.3 Segregation of Duties control
# =============================================================================

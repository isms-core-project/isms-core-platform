#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# =============================================================================
# SPDX-License-Identifier: AGPL-3.0-or-later OR LicenseRef-ISMS-Commercial
# Copyright (c) 2025-2026 ISMS Core Contributors
# =============================================================================
"""
================================================================================
ISMS-IMP-A.5.3.3 - Role-Function Mapping
================================================================================

ISO/IEC 27001:2022 Control A.5.3: Segregation of Duties
Assessment Workbook 3 of 4: Role-Function Mapping

This script generates an Excel workbook for documenting the mapping between
organisational roles, application roles, functions, and permissions.
================================================================================
"""

import logging
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# =============================================================================
# LOGGING CONFIGURATION
# =============================================================================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

# =============================================================================
# DOCUMENT METADATA
# =============================================================================
DOCUMENT_ID = "ISMS-IMP-A.5.3.3"
WORKBOOK_NAME = "Role Function Mapping"
CONTROL_ID = "A.5.3"
CONTROL_NAME = "Segregation of Duties"
CONTROL_REF = f"ISO/IEC 27001:2022 - Control {CONTROL_ID}: {CONTROL_NAME}"

# Timestamps
GENERATED_DATE = datetime.now().strftime("%d.%m.%Y")
GENERATED_TIMESTAMP = datetime.now().strftime("%Y%m%d")

# Output filename
OUTPUT_FILENAME = f"{DOCUMENT_ID}_{WORKBOOK_NAME.replace(' ', '_')}_{GENERATED_TIMESTAMP}.xlsx"

# =============================================================================
# STYLING CONSTANTS
# =============================================================================
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
TITLE_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

INPUT_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
FORMULA_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

SOD_SENSITIVE_FILL = PatternFill(start_color="FABF8F", end_color="FABF8F", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# =============================================================================
# DATA CONSTANTS
# =============================================================================
PROCESS_DOMAINS = [
    "Financial",
    "IT Operations",
    "HR",
    "Procurement",
    "Security",
    "Change Management",
    "Other"
]

RISK_LEVELS = ["Critical", "High", "Medium", "Low"]

DEPARTMENTS = [
    "Executive",
    "Finance",
    "IT",
    "Operations",
    "HR",
    "Legal",
    "Sales",
    "Marketing",
    "Engineering",
    "Support",
    "Procurement",
    "Security"
]

ROLE_TYPES = ["Composite", "Single"]

FUNCTION_CATEGORIES = [
    "Create",
    "Read",
    "Update",
    "Delete",
    "Approve",
    "Execute",
    "Admin"
]

PERMISSION_TYPES = [
    "Transaction",
    "Report",
    "API",
    "Config",
    "Data"
]

GRANT_TYPES = [
    "Direct",
    "Inherited",
    "Delegated",
    "Emergency"
]

CONFLICT_TYPES = ["X", "C", "M"]

VALIDATION_STATUSES = [
    "Validated",
    "Requires Investigation",
    "Remediated",
    "Deferred"
]

CHANGE_TYPES = [
    "Permission Added",
    "Permission Removed",
    "Function Modified",
    "Role Created",
    "Role Deleted"
]

REVIEW_FREQUENCIES = ["Monthly", "Quarterly", "Semi-Annual", "Annual"]


# =============================================================================
# WORKBOOK GENERATION FUNCTIONS
# =============================================================================

def apply_header_style(cell):
    """Apply standard header styling to a cell."""
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGNMENT
    cell.border = THIN_BORDER


def apply_input_style(cell):
    """Apply input cell styling."""
    cell.fill = INPUT_FILL
    cell.border = THIN_BORDER


def create_instructions_sheet(ws):
    """Create the Instructions sheet."""
    ws.title = "Instructions"

    instructions = [
        [f"{DOCUMENT_ID} - {WORKBOOK_NAME}"],
        [""],
        ["PURPOSE"],
        ["This workbook maps organisational roles to application roles, functions, and permissions."],
        ["It enables function-level conflict detection for segregation of duties analysis."],
        [""],
        ["SHEETS"],
        ["1. Instructions - This guidance sheet"],
        ["2. Business_Roles - Organisational role definitions"],
        ["3. Application_Roles - System-specific role inventory"],
        ["4. Functions - Discrete capability definitions"],
        ["5. Permissions - Technical permission details"],
        ["6. Role_Function_Map - Role-to-function relationships"],
        ["7. Function_Conflicts - Function-level conflicts"],
        ["8. Validation_Status - RBAC validation tracking"],
        ["9. Change_Log - Permission change history"],
        [""],
        ["MAPPING HIERARCHY"],
        ["Business Role -> Application Role -> Function -> Permission"],
        [""],
        ["FUNCTION CATEGORIES (CRUD+A)"],
        ["Create - Ability to create new records"],
        ["Read - Ability to view information"],
        ["Update - Ability to modify records"],
        ["Delete - Ability to remove records"],
        ["Approve - Ability to authorise actions"],
        ["Execute - Ability to perform transactions"],
        ["Admin - System administration capability"],
        [""],
        [f"Generated: {GENERATED_DATE}"],
        [f"Control Reference: {CONTROL_REF}"],
    ]

    for row_num, row_data in enumerate(instructions, 1):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            if row_num == 1:
                cell.font = TITLE_FONT
                cell.fill = TITLE_FILL
            elif value in ["PURPOSE", "SHEETS", "MAPPING HIERARCHY", "FUNCTION CATEGORIES (CRUD+A)"]:
                cell.font = Font(bold=True, size=11)

    ws.column_dimensions['A'].width = 80


def create_business_roles_sheet(ws):
    """Create the Business Roles sheet."""
    ws.title = "Business_Roles"

    headers = [
        "Business_Role_ID", "Role_Name", "Department", "Process_Domain",
        "Role_Owner", "Description", "Risk_Level", "Last_Reviewed"
    ]

    widths = [18, 30, 20, 20, 25, 50, 12, 15]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    dept_dv = DataValidation(type="list", formula1=f'"{",".join(DEPARTMENTS)}"')
    ws.add_data_validation(dept_dv)
    dept_dv.add('C2:C200')

    domain_dv = DataValidation(type="list", formula1=f'"{",".join(PROCESS_DOMAINS)}"')
    ws.add_data_validation(domain_dv)
    domain_dv.add('D2:D200')

    risk_dv = DataValidation(type="list", formula1=f'"{",".join(RISK_LEVELS)}"')
    ws.add_data_validation(risk_dv)
    risk_dv.add('G2:G200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Sample data
    sample_roles = [
        ("BROLE-FIN-001", "Accounts Payable Manager", "Finance", "Financial",
         "CFO", "Manages AP team and payment approvals", "High", "01.01.2026"),
        ("BROLE-IT-001", "Software Developer", "IT", "IT Operations",
         "IT Director", "Develops and maintains application code", "Medium", "01.01.2026"),
        ("BROLE-IT-002", "Release Manager", "IT", "Change Management",
         "IT Director", "Manages production deployments", "High", "01.01.2026"),
    ]

    for row_idx, role in enumerate(sample_roles, 2):
        for col_idx, value in enumerate(role, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_application_roles_sheet(ws):
    """Create the Application Roles sheet."""
    ws.title = "Application_Roles"

    headers = [
        "App_Role_ID", "Application", "Role_Name", "Role_Type",
        "Description", "Business_Roles", "Criticality", "Review_Frequency"
    ]

    widths = [18, 20, 30, 15, 40, 35, 12, 15]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    type_dv = DataValidation(type="list", formula1=f'"{",".join(ROLE_TYPES)}"')
    ws.add_data_validation(type_dv)
    type_dv.add('D2:D200')

    crit_dv = DataValidation(type="list", formula1=f'"{",".join(RISK_LEVELS)}"')
    ws.add_data_validation(crit_dv)
    crit_dv.add('G2:G200')

    freq_dv = DataValidation(type="list", formula1=f'"{",".join(REVIEW_FREQUENCIES)}"')
    ws.add_data_validation(freq_dv)
    freq_dv.add('H2:H200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Sample data
    sample_roles = [
        ("AROLE-SAP-FI-001", "SAP ERP", "FI-AP-Manager", "Composite",
         "Full AP processing and approval", "BROLE-FIN-001", "High", "Quarterly"),
        ("AROLE-GIT-001", "GitHub", "Developer", "Single",
         "Code commit and branch access", "BROLE-IT-001", "Medium", "Quarterly"),
        ("AROLE-CICD-001", "Jenkins", "Release-Manager", "Single",
         "Production deployment access", "BROLE-IT-002", "High", "Monthly"),
    ]

    for row_idx, role in enumerate(sample_roles, 2):
        for col_idx, value in enumerate(role, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_functions_sheet(ws):
    """Create the Functions sheet."""
    ws.title = "Functions"

    headers = [
        "Function_ID", "Function_Name", "Category", "Application",
        "Process", "Description", "Risk_Level", "SoD_Sensitive"
    ]

    widths = [18, 30, 12, 20, 25, 45, 12, 12]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    cat_dv = DataValidation(type="list", formula1=f'"{",".join(FUNCTION_CATEGORIES)}"')
    ws.add_data_validation(cat_dv)
    cat_dv.add('C2:C200')

    risk_dv = DataValidation(type="list", formula1=f'"{",".join(RISK_LEVELS)}"')
    ws.add_data_validation(risk_dv)
    risk_dv.add('G2:G200')

    sod_dv = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(sod_dv)
    sod_dv.add('H2:H200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Sample data
    sample_functions = [
        ("FUNC-FIN-001", "Create Vendor", "Create", "SAP ERP",
         "Accounts Payable", "Create new vendor master record", "High", "Yes"),
        ("FUNC-FIN-002", "Approve Payment", "Approve", "SAP ERP",
         "Accounts Payable", "Authorise payment execution", "Critical", "Yes"),
        ("FUNC-IT-001", "Commit Code", "Create", "GitHub",
         "Software Development", "Push code changes to repository", "Medium", "Yes"),
        ("FUNC-IT-002", "Deploy Production", "Execute", "Jenkins",
         "Change Management", "Deploy code to production environment", "Critical", "Yes"),
    ]

    for row_idx, func in enumerate(sample_functions, 2):
        for col_idx, value in enumerate(func, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_permissions_sheet(ws):
    """Create the Permissions sheet."""
    ws.title = "Permissions"

    headers = [
        "Permission_ID", "Function_ID", "Application", "Permission_Name",
        "Permission_Type", "Description", "Data_Scope", "Special_Conditions"
    ]

    widths = [18, 18, 20, 25, 15, 40, 25, 35]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    type_dv = DataValidation(type="list", formula1=f'"{",".join(PERMISSION_TYPES)}"')
    ws.add_data_validation(type_dv)
    type_dv.add('E2:E200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Sample data
    sample_permissions = [
        ("PERM-SAP-001", "FUNC-FIN-001", "SAP ERP", "T-Code FK01",
         "Transaction", "Create vendor master", "Company Code 1000", ""),
        ("PERM-SAP-002", "FUNC-FIN-002", "SAP ERP", "T-Code F110",
         "Transaction", "Payment program execution", "Company Code 1000", "Limit CHF 50'000"),
        ("PERM-GIT-001", "FUNC-IT-001", "GitHub", "repo:write",
         "API", "Push to repository", "Development branches", "Requires PR"),
        ("PERM-JNK-001", "FUNC-IT-002", "Jenkins", "build:deploy",
         "Config", "Run production deployment", "Production environment", "Requires approval"),
    ]

    for row_idx, perm in enumerate(sample_permissions, 2):
        for col_idx, value in enumerate(perm, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_role_function_map_sheet(ws):
    """Create the Role-Function Mapping sheet."""
    ws.title = "Role_Function_Map"

    headers = [
        "Mapping_ID", "Business_Role_ID", "App_Role_ID", "Function_ID",
        "Grant_Type", "Justification", "Effective_Date", "Expiry_Date"
    ]

    widths = [15, 18, 18, 18, 12, 40, 15, 15]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    grant_dv = DataValidation(type="list", formula1=f'"{",".join(GRANT_TYPES)}"')
    ws.add_data_validation(grant_dv)
    grant_dv.add('E2:E200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Sample data
    sample_mappings = [
        ("MAP-001", "BROLE-FIN-001", "AROLE-SAP-FI-001", "FUNC-FIN-001",
         "Direct", "Core AP manager duty", "01.01.2026", ""),
        ("MAP-002", "BROLE-FIN-001", "AROLE-SAP-FI-001", "FUNC-FIN-002",
         "Direct", "Payment approval authority", "01.01.2026", ""),
        ("MAP-003", "BROLE-IT-001", "AROLE-GIT-001", "FUNC-IT-001",
         "Direct", "Development responsibility", "01.01.2026", ""),
        ("MAP-004", "BROLE-IT-002", "AROLE-CICD-001", "FUNC-IT-002",
         "Direct", "Release management duty", "01.01.2026", ""),
    ]

    for row_idx, mapping in enumerate(sample_mappings, 2):
        for col_idx, value in enumerate(mapping, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_function_conflicts_sheet(ws):
    """Create the Function Conflicts sheet."""
    ws.title = "Function_Conflicts"

    headers = [
        "Conflict_ID", "Function_A", "Function_B", "Conflict_Type",
        "Risk_Level", "Justification", "Mitigation"
    ]

    widths = [15, 18, 18, 12, 12, 50, 40]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    type_dv = DataValidation(type="list", formula1=f'"{",".join(CONFLICT_TYPES)}"')
    ws.add_data_validation(type_dv)
    type_dv.add('D2:D200')

    risk_dv = DataValidation(type="list", formula1=f'"{",".join(RISK_LEVELS)}"')
    ws.add_data_validation(risk_dv)
    risk_dv.add('E2:E200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))

    # Sample data
    sample_conflicts = [
        ("FCON-001", "FUNC-FIN-001", "FUNC-FIN-002", "X",
         "Critical", "Create vendor + approve payment = fictitious vendor fraud",
         "Must be separate individuals"),
        ("FCON-002", "FUNC-IT-001", "FUNC-IT-002", "X",
         "Critical", "Commit code + deploy production = malicious code deployment",
         "Require code review and separate deployer"),
    ]

    for row_idx, conflict in enumerate(sample_conflicts, 2):
        for col_idx, value in enumerate(conflict, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_validation_status_sheet(ws):
    """Create the Validation Status sheet."""
    ws.title = "Validation_Status"

    headers = [
        "Validation_ID", "Role_ID", "Validation_Date", "Validator",
        "Documented_Functions", "Actual_Functions", "Discrepancies",
        "Status", "Resolution"
    ]

    widths = [15, 18, 15, 25, 20, 18, 15, 20, 40]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    status_dv = DataValidation(type="list", formula1=f'"{",".join(VALIDATION_STATUSES)}"')
    ws.add_data_validation(status_dv)
    status_dv.add('H2:H200')

    # Format input rows and add formulas
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            if col == 7:  # Discrepancies formula
                cell.value = f"=ABS(E{row}-F{row})"
                cell.fill = FORMULA_FILL
            else:
                apply_input_style(cell)
            cell.border = THIN_BORDER


def create_change_log_sheet(ws):
    """Create the Change Log sheet."""
    ws.title = "Change_Log"

    headers = [
        "Change_ID", "Change_Date", "Role_ID", "Change_Type",
        "Description", "Requested_By", "Approved_By", "Ticket_Reference"
    ]

    widths = [15, 15, 18, 18, 50, 25, 25, 20]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

    # Data validations
    type_dv = DataValidation(type="list", formula1=f'"{",".join(CHANGE_TYPES)}"')
    ws.add_data_validation(type_dv)
    type_dv.add('D2:D200')

    # Format input rows
    for row in range(2, 51):
        for col in range(1, len(headers) + 1):
            apply_input_style(ws.cell(row=row, column=col))


def generate_workbook():
    """Generate the complete assessment workbook."""
    logger.info(f"Generating {DOCUMENT_ID} - {WORKBOOK_NAME}")

    wb = Workbook()

    # Remove default sheet
    default_sheet = wb.active

    # Create all sheets
    create_instructions_sheet(wb.create_sheet())
    create_business_roles_sheet(wb.create_sheet())
    create_application_roles_sheet(wb.create_sheet())
    create_functions_sheet(wb.create_sheet())
    create_permissions_sheet(wb.create_sheet())
    create_role_function_map_sheet(wb.create_sheet())
    create_function_conflicts_sheet(wb.create_sheet())
    create_validation_status_sheet(wb.create_sheet())
    create_change_log_sheet(wb.create_sheet())

    # Remove default sheet
    wb.remove(default_sheet)

    # Save workbook
    wb.save(OUTPUT_FILENAME)
    logger.info(f"Workbook saved: {OUTPUT_FILENAME}")

    return OUTPUT_FILENAME


# =============================================================================
# MAIN EXECUTION
# =============================================================================
if __name__ == "__main__":
    generate_workbook()

# =============================================================================
# QA_VERIFIED: 2026-02-03
# QA_STATUS: PASSED - INITIAL CREATION
# QA_TOOL: Claude Code
# CHANGES: Initial creation for A.5.3 Segregation of Duties control
# =============================================================================
